"""Generate the Training Center document pack: quick reference cards (PDF),
onboarding checklist (PDF), voiceover scripts (PDF), and PPTX slide decks.
Outputs to /app/backend/static/training/docs/."""
import sys
from pathlib import Path

sys.path.insert(0, "/app/training_production")
from scenes import VIDEOS  # noqa: E402

from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.units import cm  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.styles import ParagraphStyle  # noqa: E402
from reportlab.platypus import (  # noqa: E402
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
)

OUT = Path("/app/backend/static/training/docs")
OUT.mkdir(parents=True, exist_ok=True)

GREEN, BLUE, INK, GREY = "#0A4A1E", "#26547C", "#1A1C1E", "#525860"

H1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=20, textColor=colors.HexColor(GREEN), spaceAfter=4)
SUB = ParagraphStyle("sub", fontName="Helvetica", fontSize=10, textColor=colors.HexColor(GREY), spaceAfter=14)
H2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor(BLUE), spaceBefore=10, spaceAfter=4)
BODY = ParagraphStyle("body", fontName="Helvetica", fontSize=9.5, textColor=colors.HexColor(INK), leading=13.5)
BULLET = ParagraphStyle("bullet", parent=BODY, leftIndent=12, bulletIndent=2, spaceAfter=2)


def brand(story, title, subtitle):
    story.append(Paragraph("SALONEHCM · TRAINING CENTER", ParagraphStyle(
        "brand", fontName="Helvetica-Bold", fontSize=8, textColor=colors.HexColor("#8B6A14"))))
    story.append(Spacer(1, 6))
    story.append(Paragraph(title, H1))
    story.append(Paragraph(subtitle, SUB))


def qrc(filename, title, subtitle, sections):
    doc = SimpleDocTemplate(str(OUT / filename), pagesize=A4,
                            leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=1.4 * cm, bottomMargin=1.4 * cm)
    story = []
    brand(story, title, subtitle)
    for heading, lines in sections:
        story.append(Paragraph(heading, H2))
        for line in lines:
            story.append(Paragraph(f"• {line}", BULLET))
    doc.build(story)
    print("wrote", filename)


qrc("qrc-getting-started.pdf", "Quick Reference — Getting Started",
    "Sign-in, navigation and finding features · keep this beside your keyboard for week one.", [
    ("Signing in", [
        "Use the work email + password from your administrator; add your 6-digit authenticator code if 2FA is on.",
        "Forgot your password? Your administrator resets it in Users &amp; Access.",
    ]),
    ("Finding your way", [
        "The sidebar shows only what your role and plan unlock.",
        "Feature Directory (2nd sidebar item) lists EVERY module with its status: Available / Requires higher plan / Admin only.",
        "Dashboard KPIs update live: headcount, payroll cost, pending leave, compliance.",
    ]),
    ("Your first week", [
        "Day 1: sign in, tour the dashboard and Feature Directory.",
        "Day 2: verify your personal details and bank/NASSIT numbers in Self-Service.",
        "Day 3–5: watch your role's training video, then pass the quiz (70%) for your certificate.",
    ]),
    ("Getting help", [
        "Training Center: videos, articles, FAQs, quizzes — public, no login needed.",
        "Access problems → your administrator. Data problems → your HR officer.",
    ])])

qrc("qrc-administrator.pdf", "Quick Reference — Administrator",
    "Users, capability flags, settings, branches and the audit trail.", [
    ("Users & access", [
        "Invite by email; the invitee sets their password via the link.",
        "Reset passwords, delete leavers, and search accounts from one screen.",
    ]),
    ("Capability flags (banknote icon)", [
        "Finance Officer — reviews and approves payroll vouchers.",
        "MoF Approver — signs payroll runs and authorizes voucher payments.",
        "Separation of duties is automatic: creator ≠ approver ≠ authorizer.",
    ]),
    ("Branches & offices (Vouchers → Branches tab)", [
        "One code per branch; optional ministry link; assign a supervisor per branch.",
        "Assign employees to branches — auto voucher generation groups by branch.",
    ]),
    ("Settings", [
        "Organisation profile, plan tier, SMS/email integrations, 2FA policy, transparency portal.",
        "Superadmin only: multi-sig MoF threshold (1–5 signatures) on the Payroll page.",
    ]),
    ("Audit log", [
        "Every approval, change and payment authorisation — searchable, permanent.",
    ])])

qrc("qrc-hr-officer.pdf", "Quick Reference — HR Officer",
    "Employees, leave, attendance and documents — the daily HR rhythm.", [
    ("Employee register", [
        "Single source of truth: grade, salary, bank, NASSIT, department.",
        "Search instantly; Add Employee guides you through required statutory fields.",
    ]),
    ("Leave", [
        "Requests arrive with balances and overlapping-absence context.",
        "Approve/reject with a comment; the decision is logged and the employee notified.",
    ]),
    ("Attendance", [
        "Daily clock-ins, absences and timesheet totals feeding payroll.",
        "Corrections allowed and audit logged — finish them BEFORE the payroll run.",
    ]),
    ("Document vault", [
        "Contracts, IDs, letters — uploaded once, stored securely, tenant-scoped.",
        "Use search + category filters to retrieve anything in seconds.",
    ]),
    ("Golden rule", [
        "Clean data today prevents payroll disputes tomorrow.",
    ])])

qrc("qrc-payroll-officer.pdf", "Quick Reference — Payroll Officer",
    "Preview → run → payslips → bank files → filing, with Sierra Leone statutory rules.", [
    ("The calculation order", [
        "Basic + allowances = gross → NASSIT employee 5% → NRA PAYE (progressive bands) → loans/deductions → net.",
        "Employer NASSIT is 10%; both appear on the payslip and NASSIT schedule.",
    ]),
    ("Running payroll", [
        "ALWAYS Preview first: check headcount and totals vs last month.",
        "Gov tenants: budget check must be green; cut-off lock respected; MoF chain signs the run.",
    ]),
    ("After the run", [
        "Payslip PDFs per employee; bulk payslip SMS; bank files in Rokel / SLCB / UBA / Ecobank formats.",
    ]),
    ("Compliance", [
        "Monthly NRA PAYE return + NASSIT schedule exports in the expected format.",
        "Deadlines tracked on the Compliance page — file before they turn red.",
    ]),
    ("Never do", [
        "Never bypass a blocked budget check without a documented, authorised override.",
    ])])

qrc("qrc-vouchers.pdf", "Quick Reference — Payroll Vouchers",
    "The centralized voucher workflow and every anti-fraud rail.", [
    ("The stages", [
        "Draft → Awaiting supervisor → Submitted → Under review → Approved → Payment authorized.",
        "Returned-for-correction possible at any pre-payment stage (reason ≥ 10 characters).",
    ]),
    ("Who does what", [
        "Branch officer/supervisor creates + submits · supervisor approves · finance officer reviews + approves · MoF/admin authorizes payment.",
    ]),
    ("The rails", [
        "One voucher per branch per period — duplicates rejected.",
        "Double-pay guard: an employee cannot be on two active vouchers in one period.",
        "Immutable after submission — corrections only via official return (revision increments).",
        "Creator can never approve own voucher; authorizer ≠ creator AND ≠ approver.",
        "Payment-authorized vouchers are permanent records; drafts are the only deletable state.",
    ]),
    ("Auto-generation", [
        "Generate-from-run splits a payroll run into one draft voucher per branch; unassigned employees are reported.",
    ])])

qrc("qrc-employee.pdf", "Quick Reference — Employee Self-Service",
    "Payslips, leave and attendance — everything you can do yourself.", [
    ("Payslips", [
        "Self-Service lists every payslip with gross, deductions and net.",
        "Open any month for the full breakdown (PAYE, NASSIT, loans) and download the PDF.",
    ]),
    ("Leave", [
        "New Request → dates + leave type → submit. Your manager is notified instantly.",
        "Watch the status change; your remaining balance is always shown.",
    ]),
    ("Attendance", [
        "Clock in/out on the Attendance page; review your own history.",
        "Something wrong? Tell HR — corrections are made and logged.",
    ]),
    ("Account", [
        "Forgot password → your administrator. Keep your bank + NASSIT details current.",
    ])])

# ---- Onboarding checklist ----
doc = SimpleDocTemplate(str(OUT / "onboarding-checklist.pdf"), pagesize=A4,
                        leftMargin=1.6 * cm, rightMargin=1.6 * cm, topMargin=1.4 * cm, bottomMargin=1.4 * cm)
story = []
brand(story, "Customer Onboarding Checklist", "The 4-week rollout plan — from kickoff to first live payroll.")
WEEKS = [
    ("Week 1 — Foundation", [
        "Confirm the tenant, plan tier and admin accounts",
        "Enable 2FA policy and review security settings",
        "Import the employee register (grades, salaries, bank, NASSIT numbers)",
        "Define departments / ministries and branches with supervisors",
        "Invite HR officers and payroll officers; grant Finance/MoF flags deliberately",
    ]),
    ("Week 2 — Configuration", [
        "Verify NRA PAYE bands and NASSIT rates for the tenant",
        "Configure allowance rules and (Gov) grade & step structure",
        "Set budget codes and period allocations (Gov)",
        "Connect SMS (Twilio) and email integrations",
        "Assign employees to branches; set the multi-sig MoF threshold",
    ]),
    ("Week 3 — Dry run", [
        "Run a payroll PREVIEW and reconcile against the last manual payroll",
        "Generate test payslips and one bank file; verify formats with the bank",
        "Create one test voucher through the full approval chain",
        "All staff watch their role's training video and pass the quiz",
    ]),
    ("Week 4 — Go live", [
        "Freeze legacy spreadsheets; announce the cut-over date",
        "Run the first live payroll with finance witnessing the approval chain",
        "Send payslip SMS / publish self-service access to all employees",
        "File the first NRA PAYE return and NASSIT schedule from the system",
        "Review the audit log and hold a 30-minute retrospective",
    ]),
]
for wk, items in WEEKS:
    story.append(Paragraph(wk, H2))
    rows = [[Paragraph("☐", BODY), Paragraph(i, BODY)] for i in items]
    t = Table(rows, colWidths=[0.8 * cm, 15.6 * cm])
    t.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
doc.build(story)
print("wrote onboarding-checklist.pdf")

# ---- Voiceover scripts ----
doc = SimpleDocTemplate(str(OUT / "voiceover-scripts.pdf"), pagesize=A4,
                        leftMargin=1.6 * cm, rightMargin=1.6 * cm, topMargin=1.4 * cm, bottomMargin=1.4 * cm)
story = []
brand(story, "Training Video Voiceover Scripts",
      "Full narration text of the 7-video SaloneHCM training series — for translators, reviewers and re-recording.")
for v in VIDEOS:
    story.append(Paragraph(v["title"], H2))
    for i, sc in enumerate(v["scenes"], 1):
        label = sc.get("chapter") or f"Scene {i}"
        story.append(Paragraph(f"<b>{i}. {label}</b> — {sc['narration']}", BODY))
        story.append(Spacer(1, 5))
    story.append(Spacer(1, 8))
doc.build(story)
print("wrote voiceover-scripts.pdf")

# ---- PPTX decks ----
from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402

G = RGBColor(0x13, 0x33, 0x26)
GOLD = RGBColor(0xD9, 0xC5, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARKINK = RGBColor(0x1A, 0x1C, 0x1E)


def deck(filename, title, subtitle, slides):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = G; bg.line.fill.background()
    tb = s.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11), Inches(2.6)).text_frame
    tb.text = title
    tb.paragraphs[0].font.size, tb.paragraphs[0].font.bold, tb.paragraphs[0].font.color.rgb = Pt(44), True, WHITE
    p = tb.add_paragraph(); p.text = subtitle; p.font.size, p.font.color.rgb = Pt(18), GOLD
    p2 = tb.add_paragraph(); p2.text = "SaloneHCM · Official training series"; p2.font.size, p2.font.color.rgb = Pt(12), RGBColor(0x8F, 0xBC, 0xA8)

    for heading, bullets in slides:
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.1))
        bar.fill.solid(); bar.fill.fore_color.rgb = G; bar.line.fill.background()
        ht = s.shapes.add_textbox(Inches(0.6), Inches(0.22), Inches(12), Inches(0.7)).text_frame
        ht.text = heading
        ht.paragraphs[0].font.size, ht.paragraphs[0].font.bold, ht.paragraphs[0].font.color.rgb = Pt(26), True, WHITE
        body = s.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4)).text_frame
        body.word_wrap = True
        for i, b in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = "•  " + b
            p.font.size, p.font.color.rgb = Pt(18), DARKINK
            p.space_after = Pt(12)

    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = G; bg.line.fill.background()
    tb = s.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11), Inches(2)).text_frame
    tb.text = "Questions?"
    tb.paragraphs[0].font.size, tb.paragraphs[0].font.bold, tb.paragraphs[0].font.color.rgb = Pt(40), True, WHITE
    p = tb.add_paragraph(); p.text = "Training Center → videos · articles · FAQs · certification quizzes"
    p.font.size, p.font.color.rgb = Pt(16), GOLD
    prs.save(str(OUT / filename))
    print("wrote", filename)


deck("deck-administrator.pptx", "Administrator Essentials", "Users, access, settings, branches & the audit trail", [
    ("Agenda", ["Roles & capability flags", "Inviting and managing users", "Organisation settings & plan tiers",
                "Branches, offices & supervisors", "The audit log", "Certification quiz"]),
    ("Roles", ["Superadmin — platform & tenants", "Admin — one organisation end-to-end",
               "Employee — self-service plus whatever flags unlock",
               "Flags: Finance Officer (voucher approval) · MoF Approver (run signing & payment authorization)"]),
    ("Users & access", ["Invite by email — invitee sets their own password", "Reset passwords in one click",
                        "Grant flags with the banknote toggle — every grant is audit logged",
                        "Remove leavers promptly: access review monthly"]),
    ("Separation of duties", ["Voucher creator can NEVER approve their own voucher",
                              "Payment authorizer ≠ creator AND ≠ finance approver",
                              "Multi-sig MoF approval: dial 1–5 required signatures (superadmin)",
                              "Plan staffing: at least 3 people with financial powers"]),
    ("Branches & offices", ["Unique code per branch · optional ministry link", "One supervisor per branch approves its vouchers",
                            "Assign employees so auto voucher generation groups correctly"]),
    ("Settings & audit", ["Profile, tier, SMS/email integrations, 2FA policy, transparency portal",
                          "Audit log: every sensitive action, searchable, permanent"]),
])

deck("deck-payroll-vouchers.pptx", "Payroll & Vouchers", "Runs, statutory rules and the centralized voucher workflow", [
    ("Agenda", ["Gross-to-net in Sierra Leone", "Preview → Run → Exports", "NRA & NASSIT filing",
                "The voucher workflow", "Anti-fraud rails", "Certification quiz"]),
    ("Gross to net", ["Basic + allowances = gross", "NASSIT employee 5% (employer 10%)",
                      "NRA PAYE on progressive monthly bands — band by band", "Loans & deductions last → net pay"]),
    ("Running payroll", ["ALWAYS preview first — reconcile vs last month", "Gov: budget check must pass; cut-off lock respected",
                         "MoF chain signs the run (up to 5 signatures)",
                         "Exports: payslip PDFs · payslip SMS · Rokel/SLCB/UBA/Ecobank bank files"]),
    ("The voucher workflow", ["Draft → Supervisor → Submitted → Review → Approved → Payment authorized",
                              "Return-for-correction with a written reason (≥10 chars)", "Revision number tracks corrections forever"]),
    ("Anti-fraud rails", ["One voucher per branch per period", "Double-pay guard across vouchers",
                          "Immutable after submission", "Creator ≠ approver ≠ authorizer", "Authorized vouchers are permanent records"]),
    ("Compliance", ["Monthly NRA PAYE return export", "NASSIT contribution schedule", "Deadlines tracked — file before they turn red"]),
])

deck("deck-employee.pptx", "Employee Self-Service", "Payslips, leave and attendance in your own hands", [
    ("Agenda", ["Signing in", "Your payslips", "Requesting leave", "Attendance", "Getting help"]),
    ("Signing in", ["Work email + password from HR", "2FA code if your organisation requires it",
                    "Forgot password? Your administrator resets it in a minute"]),
    ("Your payslips", ["Every month listed with gross, deductions, net", "Open for the full PAYE/NASSIT/loan breakdown",
                       "Download the PDF for your records or bank applications"]),
    ("Leave", ["New Request → dates + type → submit", "Manager notified instantly · status visible to you",
               "Your remaining balance always shows"]),
    ("Attendance", ["Clock in/out on the Attendance page", "Review your own history",
                    "Errors? Tell HR — corrections are logged"]),
    ("Help", ["Training Center: videos, reference cards, FAQs", "Take the Self-Service quiz for your certificate"]),
])

print("ALL DOCS DONE")
